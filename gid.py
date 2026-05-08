"""
This script processes images, generating short and long descriptions via the OpenAI API.
It can process either a single image or a folder of images.

Single image mode:
  Outputs short and long descriptions directly to standard output.

Folder mode:
  Processes images in a specified folder, storing metadata in a TSV, and optionally
  copying them into a 'Described' subfolder.

  We output six columns in the TSV:
    1) OriginalFilename
    2) ShortDescription
    3) LongDescription
    4) Context
    5) Composite
    6) SHA1

Features:
  - SHA-1-based checks for previously processed images
  - Optional no-copy behavior via --no-copy
  - Control temperature and max tokens via CLI flags
  - Preserves alphabetical order of files as typically seen in Explorer/Finder
  - Shows simple progress '1 of N'
"""

import sys
import os
import re
import base64
import shutil
import hashlib
import argparse
import logging
import json
import csv
import copy
import tempfile
from pathlib import Path
from typing import Tuple, List, Dict, Optional, Any, Set
from dataclasses import dataclass
from concurrent.futures import ThreadPoolExecutor, as_completed
from threading import Lock

try:
    from openai import OpenAI
except ImportError:
    OpenAI = None

# Configure logging - simple output with just the message
logging.basicConfig(
    level=logging.INFO,
    format='%(message)s'
)
logger = logging.getLogger('gid')

# Set OpenAI HTTP request logging to WARNING level by default (only errors shown)
logging.getLogger("openai").setLevel(logging.WARNING)
logging.getLogger("httpx").setLevel(logging.WARNING)

# Valid image extensions
VALID_EXTENSIONS = (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tiff", ".webp")

DEFAULT_MODEL = "gpt-5.5"
REASONING_EFFORT_VALUES = ("none", "low", "medium", "high", "xhigh")
DEFAULT_REASONING_EFFORT = "medium"


def render_prompt_template(
    template: str,
    short_description_max_words: int,
    context: Optional[str] = None
) -> str:
    """Render supported config prompt placeholders."""
    rendered = template.replace(
        "{short_description_max_words}",
        str(short_description_max_words)
    )
    if context is not None:
        rendered = rendered.replace("{context}", context)
    return rendered


class Config:
    """Class to handle configuration from config.json, environment variables, and CLI args."""

    DEFAULT_CONFIG: Dict[str, Any] = {
        "api": {
            "model": DEFAULT_MODEL
        },
        "parameters": {
            "temperature": 1.0,
            "max_tokens": 4000,
            "reasoning_effort": DEFAULT_REASONING_EFFORT
        },
        "processing": {
            "no_copy": False,
            "no_composites": True,
            "max_workers": 0,
            "verbose": False
        },
        "output": {
            "output_folder_name": "Described",
            "tsv_filename": "descriptions.tsv"
        },
        "prompt": {
            "system_prompt": "default",
            "instructions_prompt": "instructions",
            "single_image_prompt": "Describe the following image using the required SHORT/LONG output format.",
            "composite_image_prompt": (
                "Describe the following images together as a single composite using the required "
                "SHORT/LONG output format."
            ),
            "context_template": (
                "Additional image facts provided by the user (treat as true and naturally incorporate that "
                "knowledge if helpful or necessary to inform the description): {context}"
            ),
            "short_description_max_words": 10
        }
    }
    SAMPLE_CONFIG_OVERRIDES: Dict[str, Any] = {
        "api": {
            "api_key": "..."
        }
    }
    REQUIRED_PROMPT_FIELDS = (
        "system_prompt",
        "instructions_prompt",
        "single_image_prompt",
        "composite_image_prompt",
        "context_template",
        "short_description_max_words"
    )
    PROMPT_FILE_FIELDS = (
        "system_prompt",
        "instructions_prompt",
        "single_image_prompt",
        "composite_image_prompt",
        "context_template"
    )
    
    @staticmethod
    def find_config_file() -> Optional[str]:
        """Find the config file in standard locations."""
        # Check current directory first
        if os.path.exists("config.json"):
            return "config.json"
        
        # Check user config directory
        user_config_dir = os.path.expanduser("~/.config/gid")
        user_config_file = os.path.join(user_config_dir, "config.json")
        if os.path.exists(user_config_file):
            return user_config_file
        
        return None

    @staticmethod
    def _load_json_config(config_path: str) -> Dict[str, Any]:
        """Load a JSON config file."""
        with open(config_path, 'r', encoding='utf-8') as f:
            return json.load(f)

    @staticmethod
    def sample_config() -> Dict[str, Any]:
        """Return a complete sample config, including non-runtime placeholders."""
        config = copy.deepcopy(Config.DEFAULT_CONFIG)
        api_config = config.get("api", {})
        config["api"] = {
            "api_key": Config.SAMPLE_CONFIG_OVERRIDES["api"]["api_key"],
            **api_config
        }
        return config

    @staticmethod
    def write_sample_config(output_path: str) -> None:
        """Write the built-in defaults as a sample config file."""
        output_dir = os.path.dirname(output_path)
        if output_dir:
            os.makedirs(output_dir, exist_ok=True)
        with open(output_path, "w", encoding="utf-8") as f:
            json.dump(Config.sample_config(), f, indent=2)
            f.write("\n")

    @staticmethod
    def user_config_dir() -> str:
        """Return the user-level GID config directory."""
        return os.path.expanduser("~/.config/gid")

    @staticmethod
    def prompt_base_dirs(target_dir: Optional[str], config_path: Optional[str]) -> List[str]:
        """Return base directories whose prompts/ folders should be searched."""
        base_dirs = []
        if target_dir:
            base_dirs.append(target_dir)
        if config_path:
            base_dirs.append(os.path.dirname(os.path.abspath(config_path)) or ".")
        base_dirs.append(str(Path(__file__).resolve().parent))
        base_dirs.append(os.getcwd())
        base_dirs.append(Config.user_config_dir())

        unique_dirs = []
        seen = set()
        for base_dir in base_dirs:
            normalized = os.path.abspath(os.path.expanduser(base_dir))
            if normalized in seen:
                continue
            seen.add(normalized)
            unique_dirs.append(normalized)
        return unique_dirs

    @staticmethod
    def prompt_dirs(target_dir: Optional[str], config_path: Optional[str]) -> List[str]:
        """Return existing prompts/ directories for the active config search scope."""
        dirs = []
        for base_dir in Config.prompt_base_dirs(target_dir, config_path):
            prompt_dir = os.path.join(base_dir, "prompts")
            if os.path.isdir(prompt_dir):
                dirs.append(prompt_dir)
        return dirs

    @staticmethod
    def _is_prompt_reference(value: str) -> bool:
        """Return True if a value looks like a prompt file name rather than inline text."""
        token = value.strip()
        if not token or "\n" in token or "\r" in token:
            return False
        return re.fullmatch(r"[A-Za-z0-9_.~/\\-]+(?:\.md)?", token) is not None

    @staticmethod
    def _prompt_candidates(reference: str, prompt_dirs: List[str]) -> List[str]:
        """Return possible prompt file paths for a reference."""
        reference = reference.strip()
        candidates = []

        def add(path: str) -> None:
            normalized = os.path.abspath(os.path.expanduser(path))
            if normalized not in candidates:
                candidates.append(normalized)

        expanded_reference = os.path.expanduser(reference)
        has_path_separator = os.sep in reference or (os.altsep is not None and os.altsep in reference)
        if os.path.isabs(expanded_reference) or has_path_separator:
            add(reference)
            if not reference.lower().endswith(".md"):
                add(f"{reference}.md")

        for prompt_dir in prompt_dirs:
            add(os.path.join(prompt_dir, reference))
            if not reference.lower().endswith(".md"):
                add(os.path.join(prompt_dir, f"{reference}.md"))
        return candidates

    @staticmethod
    def _read_prompt_file(prompt_path: str) -> str:
        """Read a prompt Markdown file."""
        with open(prompt_path, "r", encoding="utf-8") as f:
            return f.read().strip()

    @staticmethod
    def _resolve_prompt_value(
        field: str,
        value: Any,
        prompt_dirs: List[str],
        required_reference_fields: Set[str]
    ) -> Any:
        """Resolve a prompt config value from prompts/<name>.md when applicable."""
        if not isinstance(value, str):
            return value
        stripped = value.strip()
        if not stripped:
            return value
        is_reference = Config._is_prompt_reference(stripped)
        if not is_reference and field not in required_reference_fields:
            return value

        for candidate in Config._prompt_candidates(stripped, prompt_dirs):
            if os.path.isfile(candidate):
                return Config._read_prompt_file(candidate)

        if field in required_reference_fields or is_reference:
            searched = ", ".join(prompt_dirs) if prompt_dirs else "no prompts directories found"
            raise FileNotFoundError(
                f"Prompt file not found for prompt.{field}: {stripped!r}. "
                f"Expected prompts/{stripped}.md in: {searched}"
            )
        return value

    @staticmethod
    def resolve_prompt_references(config: Dict[str, Any]) -> None:
        """Resolve prompt field file references in-place."""
        prompt_config = config.get("prompt")
        if not isinstance(prompt_config, dict):
            return
        metadata = config.setdefault("_gid", {})
        if metadata.get("prompts_resolved"):
            return
        prompt_dirs = metadata.setdefault("prompt_dirs", Config.prompt_dirs(None, None))
        required_reference_fields = set(metadata.setdefault("required_prompt_reference_fields", []))
        for field in Config.PROMPT_FILE_FIELDS:
            if field not in prompt_config:
                continue
            prompt_config[field] = Config._resolve_prompt_value(
                field,
                prompt_config[field],
                prompt_dirs,
                required_reference_fields
            )
        metadata["prompts_resolved"] = True

    @staticmethod
    def load_config(config_path: Optional[str] = None, require_exists: bool = False) -> Dict[str, Any]:
        """Load configuration from a JSON file."""
        config = copy.deepcopy(Config.DEFAULT_CONFIG)

        # If no config path provided, try to find one
        if not config_path:
            config_path = Config.find_config_file()
        
        # If we found a config file, load and merge it
        if config_path:
            if not os.path.exists(config_path):
                if require_exists:
                    raise FileNotFoundError(f"Config file not found: {config_path}")
                return config
            try:
                user_config = Config._load_json_config(config_path)
                
                # Deep merge the user config into the default config
                Config._merge_configs(config, user_config)

                logger.info(f"Loaded configuration from {os.path.abspath(config_path)}")
            except Exception as e:
                if require_exists:
                    raise ValueError(f"Error loading config from {config_path}: {str(e)}") from e
                logger.warning(f"Error loading config from {config_path}: {str(e)}")
        
        return config
    
    @staticmethod
    def _merge_configs(target: Dict[str, Any], source: Dict[str, Any]) -> None:
        """Recursively merge source config into target config."""
        for key, value in source.items():
            if key in target and isinstance(target[key], dict) and isinstance(value, dict):
                Config._merge_configs(target[key], value)
            else:
                target[key] = value

    @staticmethod
    def validate_prompt_config(config: Dict[str, Any]) -> None:
        """Ensure API prompt settings come from config."""
        Config.resolve_prompt_references(config)
        prompt_config = config.get("prompt")
        if not isinstance(prompt_config, dict):
            raise ValueError(
                "Missing prompt configuration. Copy config.json.sample to config.json "
                "or provide a config file with a prompt section."
            )

        missing = [
            field
            for field in Config.REQUIRED_PROMPT_FIELDS
            if field not in prompt_config or prompt_config[field] in (None, "")
        ]
        if missing:
            missing_fields = ", ".join(missing)
            raise ValueError(
                f"Missing prompt configuration value(s): {missing_fields}. "
                "Use config.json.sample as a reference."
            )

        max_words = prompt_config["short_description_max_words"]
        if not isinstance(max_words, int) or max_words < 1:
            raise ValueError("prompt.short_description_max_words must be an integer of at least 1.")

    @staticmethod
    def normalize_model_name(model_name: Any) -> str:
        """Validate and normalize a model ID before sending it to OpenAI."""
        if not isinstance(model_name, str) or not model_name.strip():
            raise ValueError("api.model must be a non-empty OpenAI model ID.")
        return model_name.strip()

    @staticmethod
    def normalize_reasoning_effort(reasoning_effort: Any) -> Optional[str]:
        """Validate and normalize a reasoning effort value."""
        if reasoning_effort is None:
            return None
        if not isinstance(reasoning_effort, str):
            raise ValueError("parameters.reasoning_effort must be a string.")
        normalized = reasoning_effort.strip().lower()
        if normalized not in REASONING_EFFORT_VALUES:
            valid_values = ", ".join(REASONING_EFFORT_VALUES)
            raise ValueError(f"parameters.reasoning_effort must be one of: {valid_values}.")
        return normalized

    @staticmethod
    def normalize_max_workers(max_workers: Any) -> Optional[int]:
        """Validate worker count. Zero means use ThreadPoolExecutor's default."""
        if max_workers is None:
            return None
        if isinstance(max_workers, bool) or not isinstance(max_workers, int):
            raise ValueError("processing.max_workers must be an integer.")
        if max_workers < 0:
            raise ValueError("processing.max_workers must be 0 for auto or at least 1.")
        return None if max_workers == 0 else max_workers


@dataclass
class ImageResult:
    """Data class to store image processing results."""
    original_filename: str
    image_path: str
    file_hash: str
    short_desc: str
    long_desc: str
    context: str


@dataclass
class TSVEntry:
    """Data class to store TSV row data."""
    original_filename: str
    short_desc: str
    long_desc: str
    context: str
    composite: bool
    file_hash: str


@dataclass
class CompositeResult:
    """Data class to store composite image processing results."""
    entry: TSVEntry
    composite_hash: str
    short_desc: str
    long_desc: str


class FileHelper:
    """Helper class for file operations."""
    
    @staticmethod
    def hash_file(file_path: str, chunk_size: int = 65536) -> str:
        """
        Compute SHA-1 hash of a file to facilitate binary comparison.
        """
        sha1 = hashlib.sha1()
        with open(file_path, "rb") as f:
            for chunk in iter(lambda: f.read(chunk_size), b""):
                sha1.update(chunk)
        return sha1.hexdigest()
    
    @staticmethod
    def encode_image(image_path: str) -> str:
        """
        Encode the given image file into base64.
        """
        with open(image_path, "rb") as image_file:
            return base64.b64encode(image_file.read()).decode("utf-8")
    
    @staticmethod
    def sanitize_filename(name: str) -> str:
        """
        Replace invalid Windows filename characters (\\/:*?"<>|) with spaces,
        and remove trailing periods/spaces.
        """
        name = re.sub(r'[\\/:*?"<>|]+', " ", name)
        name = re.sub(r"\s+", " ", name).strip()
        name = name.rstrip(" .")
        return name or "image"
    
    @staticmethod
    def described_folder_path(folder_path: str, output_folder_name: str, no_copy: bool = False) -> str:
        """
        Return the folder where TSV output and copied files belong.

        Args:
            folder_path: Source folder path
            output_folder_name: Name of output subfolder
            no_copy: If True, return source folder directly

        Returns:
            Path where TSV and copied files should be placed
        """
        if no_copy:
            return folder_path
        return os.path.join(folder_path, output_folder_name)

    @staticmethod
    def ensure_folder(folder_path: str) -> None:
        """Create a folder if a write operation needs it."""
        os.makedirs(folder_path, exist_ok=True)


class TSVHandler:
    """Class to handle TSV file operations."""

    SMART_PUNCTUATION_TRANSLATION = str.maketrans({
        "\u2018": "'",
        "\u2019": "'",
        "\u201A": "'",
        "\u201B": "'",
        "\u201C": '"',
        "\u201D": '"',
        "\u201E": '"',
        "\u201F": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u2212": "-",
        "\u2026": "...",
        "\u00A0": " "
    })
    
    def __init__(self, tsv_path: str):
        self.tsv_path = tsv_path
        self.entries: List[TSVEntry] = []
        self.entries_by_hash: Dict[str, TSVEntry] = {}
        self.has_context_column = False
        self.has_composite_column = False
        self.load()
    
    @staticmethod
    def _escape_newlines(text: str) -> str:
        """Escape actual newlines as \\n for single-line TSV rows."""
        text = text.translate(TSVHandler.SMART_PUNCTUATION_TRANSLATION)
        return text.replace('\n', '\\n').replace('\r', '\\r')

    @staticmethod
    def _unescape_newlines(text: str) -> str:
        """Unescape \\n back to actual newlines when reading."""
        return text.replace('\\n', '\n').replace('\\r', '\r')

    def load(self) -> None:
        """
        Load existing lines from the TSV (skipping header).
        """
        if not os.path.exists(self.tsv_path):
            return
        
        with open(self.tsv_path, "r", encoding="utf-8-sig", newline="") as tsv_file:
            reader = csv.reader(tsv_file, delimiter="\t")
            try:
                first_row = next(reader)
            except StopIteration:
                return

            expected_columns = {
                "OriginalFilename",
                "ShortDescription",
                "LongDescription",
                "Context",
                "Composite",
                "SHA1"
            }
            if any(col in expected_columns for col in first_row):
                header_cols = first_row
                data_rows = reader
            else:
                header_cols = []
                data_rows = [first_row] + list(reader)

            col_index = {col: idx for idx, col in enumerate(header_cols)}
            if header_cols:
                self.has_context_column = "Context" in col_index
                self.has_composite_column = "Composite" in col_index
            for parts in data_rows:
                if not parts or not any(part.strip() for part in parts):
                    continue
                if not header_cols:
                    if len(parts) >= 6:
                        orig_filename, short_desc, long_desc, context, composite_raw, hash_value = parts[:6]
                    elif len(parts) >= 5:
                        orig_filename, short_desc, long_desc, context, hash_value = parts[:5]
                        composite_raw = ""
                    elif len(parts) >= 4:
                        orig_filename, short_desc, long_desc, hash_value = parts[:4]
                        context = ""
                        composite_raw = ""
                    else:
                        continue
                else:
                    def col(name: str, default: str = "") -> str:
                        idx = col_index.get(name)
                        if idx is None or idx >= len(parts):
                            return default
                        return parts[idx]

                    orig_filename = col("OriginalFilename", parts[0] if parts else "")
                    short_desc = col("ShortDescription", parts[1] if len(parts) > 1 else "")
                    long_desc = col("LongDescription", parts[2] if len(parts) > 2 else "")
                    context = col("Context", "")
                    composite_raw = col("Composite", "")
                    hash_value = col("SHA1", parts[-1] if parts else "")

                composite = composite_raw.strip().lower() == "yes"
                entry = TSVEntry(
                    original_filename=orig_filename,
                    short_desc=self._unescape_newlines(short_desc),
                    long_desc=self._unescape_newlines(long_desc),
                    context=self._unescape_newlines(context),
                    composite=composite,
                    file_hash=hash_value
                )
                self.entries.append(entry)
                if hash_value and hash_value not in self.entries_by_hash:
                    self.entries_by_hash[hash_value] = entry

    def get_entry(self, file_hash: str) -> Optional[TSVEntry]:
        """Get an entry by hash if it exists."""
        return self.entries_by_hash.get(file_hash)

    def get_composite_entries(self) -> List[TSVEntry]:
        """Return all composite entries."""
        return [entry for entry in self.entries if entry.composite]

    def update_entry_hash(self, entry: TSVEntry, new_hash: str) -> None:
        """Update the hash for an entry and maintain the lookup map."""
        old_hash = entry.file_hash
        if old_hash and self.entries_by_hash.get(old_hash) is entry:
            del self.entries_by_hash[old_hash]
        entry.file_hash = new_hash
        if new_hash and new_hash not in self.entries_by_hash:
            self.entries_by_hash[new_hash] = entry

    def upsert_entry(
        self,
        orig_filename: str,
        short_desc: str,
        long_desc: str,
        context: str,
        file_hash: str,
        composite: bool = False
    ) -> None:
        """Insert or update an entry keyed by file hash."""
        entry = self.entries_by_hash.get(file_hash)
        if entry is None:
            entry = TSVEntry(
                original_filename=orig_filename,
                short_desc=short_desc,
                long_desc=long_desc,
                context=context,
                composite=composite,
                file_hash=file_hash
            )
            self.entries.append(entry)
            if file_hash:
                self.entries_by_hash[file_hash] = entry
            return

        if not entry.original_filename:
            entry.original_filename = orig_filename
        if short_desc:
            entry.short_desc = short_desc
        if long_desc:
            entry.long_desc = long_desc
        if context:
            entry.context = context
        entry.composite = entry.composite or composite
        if file_hash and file_hash not in self.entries_by_hash:
            self.entries_by_hash[file_hash] = entry

    def write_all(self) -> None:
        """Rewrite the TSV with the current entries."""
        folder = os.path.dirname(self.tsv_path) or "."
        FileHelper.ensure_folder(folder)
        fd, temp_path = tempfile.mkstemp(
            prefix=f".{os.path.basename(self.tsv_path)}.",
            suffix=".tmp",
            dir=folder,
            text=True
        )
        try:
            with os.fdopen(fd, "w", encoding="utf-8", newline="") as tsv_file:
                writer = csv.writer(tsv_file, delimiter="\t", lineterminator="\n")
                writer.writerow([
                    "OriginalFilename",
                    "ShortDescription",
                    "LongDescription",
                    "Context",
                    "Composite",
                    "SHA1"
                ])
                for entry in self.entries:
                    escaped_short = self._escape_newlines(entry.short_desc)
                    escaped_long = self._escape_newlines(entry.long_desc)
                    escaped_context = self._escape_newlines(entry.context)
                    composite_value = "yes" if entry.composite else "no"
                    writer.writerow([
                        entry.original_filename,
                        escaped_short,
                        escaped_long,
                        escaped_context,
                        composite_value,
                        entry.file_hash
                    ])
            os.replace(temp_path, self.tsv_path)
        except Exception:
            try:
                os.unlink(temp_path)
            except OSError:
                pass
            raise

    def write_excel(self, xlsx_path: str) -> bool:
        """Write the current entries to an Excel .xlsx file."""
        try:
            from openpyxl import Workbook
        except ImportError:
            logger.error("openpyxl is required for --make-excel. Install with: pip install openpyxl")
            return False

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "descriptions"
        worksheet.append([
            "OriginalFilename",
            "ShortDescription",
            "LongDescription",
            "Context",
            "Composite",
            "SHA1"
        ])
        for entry in self.entries:
            composite_value = "yes" if entry.composite else "no"
            worksheet.append([
                entry.original_filename,
                entry.short_desc,
                entry.long_desc,
                entry.context,
                composite_value,
                entry.file_hash
            ])

        workbook.save(xlsx_path)
        return True


class ImageDescriber:
    """Class to interact with OpenAI API for image description."""

    DANGLING_SHORT_END_WORDS = {
        "a",
        "an",
        "and",
        "as",
        "at",
        "beside",
        "between",
        "by",
        "featuring",
        "for",
        "from",
        "in",
        "inside",
        "into",
        "near",
        "of",
        "on",
        "or",
        "outside",
        "over",
        "showing",
        "the",
        "through",
        "titled",
        "to",
        "under",
        "wearing",
        "while",
        "with"
    }
    GENERIC_LONG_DESCRIPTION_OPENING_PATTERN = re.compile(
        r"^(?:the|this|an?)\s+"
        r"(?:image|photo|photograph|picture|screenshot|screen\s+capture|visual)\s+"
        r"(?:shows|depicts|features|presents|displays|captures|contains|includes)\b",
        re.IGNORECASE
    )
    
    def __init__(
        self,
        api_key: str,
        model: str = DEFAULT_MODEL,
        temperature: float = 1.0,
        max_tokens: int = 4000,
        reasoning_effort: Optional[str] = DEFAULT_REASONING_EFFORT,
        system_prompt: Optional[str] = None,
        instructions_prompt: Optional[str] = None,
        single_image_prompt: Optional[str] = None,
        composite_image_prompt: Optional[str] = None,
        context_template: Optional[str] = None,
        short_description_max_words: Optional[int] = None
    ):
        if OpenAI is None:
            raise ImportError(
                "The openai package is required for image description. "
                "Install it with: pip install openai"
            )
        self.api_key = api_key
        self.model = Config.normalize_model_name(model)
        self.temperature = temperature
        self.max_tokens = max_tokens
        self.reasoning_effort = Config.normalize_reasoning_effort(reasoning_effort)
        self.client = OpenAI(api_key=api_key)
        if not system_prompt:
            raise ValueError("Missing prompt configuration value: system_prompt.")
        if not instructions_prompt:
            raise ValueError("Missing prompt configuration value: instructions_prompt.")
        if not single_image_prompt:
            raise ValueError("Missing prompt configuration value: single_image_prompt.")
        if not composite_image_prompt:
            raise ValueError("Missing prompt configuration value: composite_image_prompt.")
        if not context_template:
            raise ValueError("Missing prompt configuration value: context_template.")
        if not isinstance(short_description_max_words, int) or short_description_max_words < 1:
            raise ValueError("prompt.short_description_max_words must be an integer of at least 1.")
        self.short_description_max_words = short_description_max_words
        combined_system_prompt = f"{instructions_prompt.strip()}\n\n{system_prompt.strip()}"
        self.system_prompt = render_prompt_template(
            combined_system_prompt,
            self.short_description_max_words
        )
        self.single_image_prompt = single_image_prompt
        self.composite_image_prompt = composite_image_prompt
        self.context_template = context_template

    @staticmethod
    def _label_pattern(label: str) -> str:
        """Return a tolerant response label pattern for SHORT/LONG labels."""
        return rf"(?m)^\s*{label}(?:\s*description|description)?\b\s*(?:[:\uFF1A]|[\-\u2013\u2014]\s+)\s*"

    @staticmethod
    def _normalize_response_text(text: str) -> str:
        """Normalize model text before parsing or validating fields."""
        return (
            text
            .replace("\r\n", "\n")
            .replace("\r", "\n")
            .replace("\\r\\n", "\n")
            .replace("\\n", "\n")
            .replace("\\r", "\n")
        )

    @staticmethod
    def _collapse_inline_whitespace(text: str) -> str:
        """Collapse all whitespace to single spaces."""
        return re.sub(r"\s+", " ", text).strip()

    @staticmethod
    def _plain_word(word: str) -> str:
        """Return a lowercased word stripped of edge punctuation."""
        return re.sub(r"^[^A-Za-z0-9]+|[^A-Za-z0-9]+$", "", word).lower()

    @classmethod
    def _trim_dangling_short_end(cls, short_desc: str) -> str:
        """Remove obvious dangling tail words created by max-word trimming."""
        short_desc = short_desc.strip().rstrip(" .,:;")
        if short_desc.count('"') % 2 == 1:
            short_desc = short_desc[:short_desc.rfind('"')].strip().rstrip(" .,:;")

        while short_desc:
            words = short_desc.split()
            if not words:
                return ""

            if len(words) >= 2:
                previous_word = cls._plain_word(words[-2])
                last_word = cls._plain_word(words[-1])
                if previous_word in cls.DANGLING_SHORT_END_WORDS and last_word.endswith("ing"):
                    short_desc = " ".join(words[:-1]).strip().rstrip(" .,:;")
                    continue

            last_word = cls._plain_word(words[-1])
            if last_word in cls.DANGLING_SHORT_END_WORDS:
                short_desc = " ".join(words[:-1]).strip().rstrip(" .,:;")
                continue
            break
        return short_desc

    def _clean_short_description(self, short_desc: str) -> str:
        """Clean a short description without inventing new content."""
        short_desc = self._normalize_response_text(short_desc)
        short_desc = re.sub(self._label_pattern("short"), "", short_desc, flags=re.IGNORECASE)
        short_desc = self._collapse_inline_whitespace(short_desc).strip('"')
        short_desc = self._limit_short_description(short_desc)
        return self._trim_dangling_short_end(short_desc)

    def _clean_long_description(self, long_desc: str) -> str:
        """Clean labels from a long description while preserving paragraph breaks."""
        long_desc = self._normalize_response_text(long_desc).strip()
        long_desc = re.sub(self._label_pattern("long"), "", long_desc, count=1, flags=re.IGNORECASE).strip()
        return long_desc

    @classmethod
    def description_fields_are_malformed(
        cls,
        short_desc: str,
        long_desc: str,
        short_description_max_words: int
    ) -> bool:
        """Return True when stored descriptions look like parser failures."""
        short_desc = cls._normalize_response_text(short_desc or "").strip()
        long_desc = cls._normalize_response_text(long_desc or "").strip()
        if not short_desc or not long_desc:
            return True

        short_label = cls._label_pattern("short")
        long_label = cls._label_pattern("long")
        if re.search(short_label, short_desc, flags=re.IGNORECASE):
            return True
        if re.search(long_label, short_desc, flags=re.IGNORECASE):
            return True
        if re.search(short_label, long_desc, flags=re.IGNORECASE):
            return True
        if re.search(long_label, long_desc, flags=re.IGNORECASE):
            return True
        if "\n" in short_desc:
            return True
        if "_" in short_desc:
            return True
        if len(short_desc.split()) > short_description_max_words:
            return True
        if cls._trim_dangling_short_end(short_desc) != short_desc.rstrip(" .,:;"):
            return True
        if cls.GENERIC_LONG_DESCRIPTION_OPENING_PATTERN.search(long_desc):
            return True
        return False

    def _parse_description_response(self, text_response: str) -> Optional[Tuple[str, str]]:
        """Parse the model response into short and long descriptions."""
        text_response = self._normalize_response_text(text_response)
        short_pattern = self._label_pattern("short")
        long_pattern = self._label_pattern("long")
        short_match = re.search(short_pattern, text_response, flags=re.IGNORECASE)
        long_match = re.search(long_pattern, text_response, flags=re.IGNORECASE)
        if short_match and long_match and long_match.start() > short_match.end():
            raw_short_part = text_response[short_match.end():long_match.start()].strip()
            if "\n" in raw_short_part:
                return None
            short_part = self._clean_short_description(raw_short_part)
            long_part = self._clean_long_description(text_response[long_match.end():])
            if (
                short_part and long_part
                and not self.description_fields_are_malformed(
                    short_part,
                    long_part,
                    self.short_description_max_words
                )
            ):
                return short_part, long_part

        if short_match or long_match:
            return None

        # Legacy two-line format: first non-empty line is short, rest is long.
        # Only accept it if the first line plausibly fits the short-description field.
        lines = [line.strip() for line in text_response.splitlines() if line.strip()]
        if len(lines) >= 2:
            first_line = lines[0]
            first_line_words = first_line.split()
            if first_line_words and len(first_line_words) <= self.short_description_max_words + 5:
                short_part = self._clean_short_description(first_line)
                long_part = self._clean_long_description("\n".join(lines[1:]))
                if (
                    short_part and long_part
                    and not self.description_fields_are_malformed(
                        short_part,
                        long_part,
                        self.short_description_max_words
                    )
                ):
                    return short_part, long_part

        return None
    
    def _limit_short_description(self, short_desc: str) -> str:
        """Trim short description to the configured max word count."""
        max_words = self.short_description_max_words
        if not max_words or max_words <= 0:
            return short_desc
        words = short_desc.split()
        if len(words) <= max_words:
            return short_desc
        return " ".join(words[:max_words])
    
    def describe_image(self, image_path: str, context: str = "") -> Tuple[str, str]:
        """Call the OpenAI API to describe a single image."""
        return self.describe_images([image_path], context)

    def describe_images(self, image_paths: List[str], context: str = "") -> Tuple[str, str]:
        """
        Call the OpenAI API to generate short and long descriptions of one or more images.
        Returns (short_desc, long_desc).
        """
        try:
            content = []
            prompt_text = self.single_image_prompt
            if len(image_paths) > 1:
                prompt_text = self.composite_image_prompt
            prompt_text = render_prompt_template(
                prompt_text,
                self.short_description_max_words
            )
            content.append({"type": "input_text", "text": prompt_text})
            for image_path in image_paths:
                encoded_image = FileHelper.encode_image(image_path)
                extension = os.path.splitext(image_path)[1][1:].lower()
                data_url = f"data:image/{extension};base64,{encoded_image}"
                content.append({"type": "input_image", "image_url": data_url})

            instructions = self.system_prompt
            if context.strip():
                context_value = context.strip()
                template = self.context_template or ""
                if "{context}" in template:
                    context_text = render_prompt_template(
                        template,
                        self.short_description_max_words,
                        context_value
                    )
                else:
                    rendered_template = render_prompt_template(
                        template,
                        self.short_description_max_words
                    )
                    context_text = f"{rendered_template} {context_value}".strip() if rendered_template else context_value
                instructions = f"{instructions.strip()}\n\n{context_text}"

            response_params = {
                "model": self.model,
                "input": [
                    {
                        "role": "user",
                        "content": content
                    }
                ],
                "instructions": instructions,
                "max_output_tokens": self.max_tokens
            }
            if self.temperature is not None and self.temperature != 1.0:
                response_params["temperature"] = self.temperature
            if self.reasoning_effort is not None:
                response_params["reasoning"] = {"effort": self.reasoning_effort}

            response = self.client.responses.create(**response_params)
            
            text_response = response.output_text
            if text_response is None:
                logger.error("Empty response from API.")
                return "Error", "Empty response from API"
            text_response = text_response.strip()

            parsed_response = self._parse_description_response(text_response)
            if parsed_response:
                return parsed_response

            logger.warning(f"Unexpected response format for {image_path}")
            return "Error", "Unexpected response format from API; expected SHORT and LONG labels."
        except Exception as e:
            logger.error(f"Error describing image {image_path}: {str(e)}")
            return "Error", f"Error: {str(e)}"


class ImageProcessor:
    """Main class to orchestrate image processing."""
    
    def __init__(self, folder_path: str, config: Dict[str, Any], init_only: bool = False):
        self.folder_path = folder_path
        self.api_key = config["api"]["api_key"]
        self.model = Config.normalize_model_name(config["api"]["model"])
        self.temperature = config["parameters"]["temperature"]
        self.max_tokens = config["parameters"]["max_tokens"]
        self.reasoning_effort = Config.normalize_reasoning_effort(
            config["parameters"].get("reasoning_effort")
        )
        self.no_copy = config["processing"]["no_copy"]
        self.no_composites = config["processing"].get("no_composites", True)
        self.max_workers = Config.normalize_max_workers(config["processing"]["max_workers"])
        self.verbose = config["processing"]["verbose"]
        self.output_folder_name = config["output"]["output_folder_name"]
        self.tsv_filename = config["output"]["tsv_filename"]
        prompt_config = config.get("prompt", {})
        self.system_prompt = prompt_config.get("system_prompt")
        self.instructions_prompt = prompt_config.get("instructions_prompt")
        self.single_image_prompt = prompt_config.get("single_image_prompt")
        self.composite_image_prompt = prompt_config.get("composite_image_prompt")
        self.context_template = prompt_config.get("context_template")
        self.short_description_max_words = prompt_config.get("short_description_max_words")
        
        self.described_folder_path = FileHelper.described_folder_path(folder_path, self.output_folder_name, self.no_copy)
        self.tsv_path = os.path.join(self.described_folder_path, self.tsv_filename)
        self.tsv_handler = TSVHandler(self.tsv_path)
        self.describer = None
        if not init_only:
            Config.validate_prompt_config(config)
            self.describer = ImageDescriber(
                api_key=self.api_key,
                model=self.model,
                temperature=self.temperature,
                max_tokens=self.max_tokens,
                reasoning_effort=self.reasoning_effort,
                system_prompt=self.system_prompt,
                instructions_prompt=self.instructions_prompt,
                single_image_prompt=self.single_image_prompt,
                composite_image_prompt=self.composite_image_prompt,
                context_template=self.context_template,
                short_description_max_words=self.short_description_max_words
            )
        self.progress_lock = Lock()
        
        # Set up logging based on verbosity
        if self.verbose:
            # Enable HTTP request logging in verbose mode
            logging.getLogger("openai").setLevel(logging.INFO)
            logging.getLogger("httpx").setLevel(logging.INFO)

    def entry_needs_description(self, entry: TSVEntry) -> bool:
        """Return True when a TSV row is missing or has malformed descriptions."""
        if not entry.short_desc or not entry.long_desc:
            return True
        return ImageDescriber.description_fields_are_malformed(
            entry.short_desc,
            entry.long_desc,
            self.short_description_max_words
        )

    def collect_image_files(
        self,
        include_existing: bool = False,
        skip_filenames: Optional[Set[str]] = None
    ) -> List[Tuple[str, str, str, str]]:
        """
        Collect image files that need processing.
        Returns a list of tuples: (filename, image_path, file_hash, context)
        """
        # Sort by name ignoring case, to mimic typical Explorer/Finder order
        all_files = sorted(os.listdir(self.folder_path), key=str.lower)
        
        tasks = []
        pending_hashes = set()
        
        # Identify which files need describing, preserving sorted order
        for filename in all_files:
            if filename.lower().endswith(VALID_EXTENSIONS):
                if skip_filenames and filename.lower() in skip_filenames:
                    continue
                image_path = os.path.join(self.folder_path, filename)
                file_hash = FileHelper.hash_file(image_path)
                entry = self.tsv_handler.get_entry(file_hash)
                if entry and not include_existing and not self.entry_needs_description(entry):
                    continue
                # Skip if queued
                if file_hash in pending_hashes:
                    continue
                pending_hashes.add(file_hash)
                context = entry.context if entry else ""
                tasks.append((filename, image_path, file_hash, context))
        
        return tasks
    
    def process_image(self, task: Tuple[str, str, str, str]) -> Optional[ImageResult]:
        """Process a single image and return the result."""
        filename, image_path, file_hash, context = task
        if not self.describer:
            logger.error("Image describer is not initialized.")
            return None
        short_desc, long_desc = self.describer.describe_image(image_path, context)
        
        if short_desc == "Error":
            return None
        
        return ImageResult(
            original_filename=filename,
            image_path=image_path,
            file_hash=file_hash,
            short_desc=short_desc,
            long_desc=long_desc,
            context=context
        )

    @staticmethod
    def _strip_image_extension(name: str) -> str:
        """Remove a known image extension from the end of the name, if present."""
        lower = name.lower()
        for ext in VALID_EXTENSIONS:
            if lower.endswith(ext):
                return name[:-len(ext)]
        return name

    @classmethod
    def _normalize_base_name(cls, name: str) -> str:
        """Normalize a composite base name for matching."""
        return cls._strip_image_extension(name).lower()

    @staticmethod
    def _looks_like_composite_sequence(indexes: List[int], has_base_image: bool = False) -> bool:
        """
        Require composite parts to look like an intentional numbered sequence.

        This avoids treating camera filenames such as IMG_3392.jpeg as composite
        parts just because they end with an underscore and digits.
        """
        if not indexes:
            return False

        unique_indexes = sorted(set(indexes))
        if len(unique_indexes) != len(indexes):
            return False

        if unique_indexes[0] == 1:
            if len(unique_indexes) >= 2 and unique_indexes == list(range(1, len(unique_indexes) + 1)):
                return True
            return has_base_image and unique_indexes == [1]

        if unique_indexes[0] == 0:
            if len(unique_indexes) >= 2 and unique_indexes == list(range(len(unique_indexes))):
                return True
            return has_base_image and unique_indexes == [0]

        return False

    def _discover_composite_sets(self) -> List[Tuple[str, List[Tuple[int, str, str, str]]]]:
        """Discover composite sets from sequential underscore-numbered filenames."""
        pattern = re.compile(
            r"^(?P<base>.+)_(?P<index>\d+)\.([A-Za-z0-9]+)$",
            re.IGNORECASE
        )
        grouped: Dict[str, Dict[str, Any]] = {}
        all_filenames = [
            filename
            for filename in os.listdir(self.folder_path)
            if filename.lower().endswith(VALID_EXTENSIONS)
        ]
        filename_lookup: Dict[str, str] = {}
        for filename in all_filenames:
            lower = filename.lower()
            if lower not in filename_lookup:
                filename_lookup[lower] = filename
            match = pattern.match(filename)
            if not match:
                continue
            base = match.group("base")
            index = int(match.group("index"))
            image_path = os.path.join(self.folder_path, filename)
            file_hash = FileHelper.hash_file(image_path)
            base_key = base.lower()
            group = grouped.get(base_key)
            if group is None:
                group = {"base": base, "files": []}
                grouped[base_key] = group
            group["files"].append((index, filename, image_path, file_hash))

        composite_sets: List[Tuple[str, List[Tuple[int, str, str, str]]]] = []
        for group in grouped.values():
            indexed_files = list(group["files"])
            if not indexed_files:
                continue
            base = group["base"]
            base_lower = base.lower()
            base_file: Optional[Tuple[int, str, str, str]] = None
            for ext in VALID_EXTENSIONS:
                candidate_lower = f"{base_lower}{ext}"
                actual = filename_lookup.get(candidate_lower)
                if not actual:
                    continue
                image_path = os.path.join(self.folder_path, actual)
                file_hash = FileHelper.hash_file(image_path)
                base_file = (-1, actual, image_path, file_hash)
                break

            indexes = [index for index, _filename, _path, _file_hash in indexed_files]
            if not self._looks_like_composite_sequence(indexes, has_base_image=base_file is not None):
                continue

            files = list(indexed_files)
            if base_file is not None:
                files.append(base_file)
            files.sort(key=lambda item: (item[0], item[1].lower()))
            composite_sets.append((group["base"], files))

        composite_sets.sort(key=lambda item: item[0].lower())
        return composite_sets

    @staticmethod
    def _preferred_composite_name(
        base_name: str,
        files: List[Tuple[int, str, str, str]]
    ) -> str:
        """Choose the composite row's OriginalFilename, including an extension."""
        for index, filename, _path, _hash in files:
            if index == -1:
                return filename
        first_name = files[0][1]
        _, ext = os.path.splitext(first_name)
        return f"{base_name}{ext}" if ext else base_name

    @staticmethod
    def _compute_composite_hash(files: List[Tuple[int, str, str, str]]) -> str:
        """Compute a hash from the ordered list of composite file hashes."""
        sha1 = hashlib.sha1()
        for _index, filename, _path, file_hash in files:
            sha1.update(f"{filename}\t{file_hash}\n".encode("utf-8"))
        return sha1.hexdigest()

    def collect_composite_tasks(
        self
    ) -> Tuple[List[Tuple[TSVEntry, List[Tuple[int, str, str, str]], str]], Set[str]]:
        """Collect composite tasks and the filenames they cover."""
        if self.no_composites:
            return [], set()

        tasks = []
        skip_filenames: Set[str] = set()
        composite_sets = self._discover_composite_sets()
        composite_entries = self.tsv_handler.get_composite_entries()
        entry_by_base = {}
        for entry in composite_entries:
            key = self._normalize_base_name(entry.original_filename)
            if key not in entry_by_base:
                entry_by_base[key] = entry
        discovered_bases: Set[str] = set()

        for base_name, files in composite_sets:
            base_key = self._normalize_base_name(base_name)
            discovered_bases.add(base_key)
            entry = entry_by_base.get(base_key)
            preferred_name = self._preferred_composite_name(base_name, files)
            if entry is None:
                entry = TSVEntry(
                    original_filename=preferred_name,
                    short_desc="",
                    long_desc="",
                    context="",
                    composite=True,
                    file_hash=""
                )
                self.tsv_handler.entries.append(entry)
            else:
                entry.composite = True
                if self._strip_image_extension(entry.original_filename) == entry.original_filename:
                    entry.original_filename = preferred_name

            for _index, filename, _path, _file_hash in files:
                skip_filenames.add(filename.lower())

            composite_hash = self._compute_composite_hash(files)
            previous_hash = entry.file_hash
            if not self.entry_needs_description(entry) and previous_hash == composite_hash:
                continue
            tasks.append((entry, files, composite_hash))

        for entry in composite_entries:
            base_key = self._normalize_base_name(entry.original_filename)
            if base_key not in discovered_bases:
                logger.warning(f"No composite files found for base name '{entry.original_filename}'.")

        return tasks, skip_filenames

    def process_composite_task(
        self,
        task: Tuple[TSVEntry, List[Tuple[int, str, str, str]], str]
    ) -> Optional[CompositeResult]:
        """Process a composite image set and return the result."""
        entry, files, composite_hash = task
        if not self.describer:
            logger.error("Image describer is not initialized.")
            return None
        image_paths = [path for _index, _filename, path, _hash in files]
        short_desc, long_desc = self.describer.describe_images(image_paths, entry.context)
        if short_desc == "Error":
            return None
        return CompositeResult(
            entry=entry,
            composite_hash=composite_hash,
            short_desc=short_desc,
            long_desc=long_desc
        )

    def discover_composites(self) -> List[Tuple[str, List[str]]]:
        """Discover composite sets from sequential underscore-numbered filenames."""
        composites: List[Tuple[str, List[str]]] = []
        if self.no_composites:
            return composites
        for base_name, files in self._discover_composite_sets():
            composites.append((base_name, [filename for _index, filename, _path, _hash in files]))
        return composites

    def show_composites(self) -> None:
        """Print composite base names and their matching files."""
        if self.no_composites:
            logger.info("Composite detection is disabled.")
            return
        composites = self.discover_composites()
        if not composites:
            logger.info("No composites found.")
            return
        for base_name, files in composites:
            logger.info(f"{base_name}:")
            for filename in files:
                logger.info(f"  - {filename}")

    def handle_composite_result(self, result: CompositeResult) -> None:
        """Update TSV entry for a composite result (no file copying)."""
        short_desc = FileHelper.sanitize_filename(result.short_desc)
        result.entry.short_desc = short_desc
        result.entry.long_desc = result.long_desc
        self.tsv_handler.update_entry_hash(result.entry, result.composite_hash)

    def handle_image_result(self, result: ImageResult) -> None:
        """
        Handle the image processing result - copy file and add to TSV.
        """
        # Sanitize short_desc for file naming
        short_desc = FileHelper.sanitize_filename(result.short_desc)

        if self.no_copy:
            self.tsv_handler.upsert_entry(
                result.original_filename,
                short_desc,
                result.long_desc,
                result.context,
                result.file_hash
            )
            return

        FileHelper.ensure_folder(self.described_folder_path)

        # Build new filename
        _, ext = os.path.splitext(result.image_path)
        base_name = short_desc
        new_image_name = f"{base_name}{ext.lower()}"
        new_image_path = os.path.join(self.described_folder_path, new_image_name)
        
        # Handle collisions - limit to 100 attempts to avoid infinite loop
        counter = 2
        max_attempts = 100
        while os.path.exists(new_image_path) and counter <= max_attempts:
            alt_name = f"{base_name} {counter}{ext.lower()}"
            alt_path = os.path.join(self.described_folder_path, alt_name)
            if not os.path.exists(alt_path):
                new_image_name = alt_name
                new_image_path = alt_path
                break
            counter += 1
        
        # If we couldn't find a non-colliding name, skip this file
        if counter > max_attempts:
            logger.warning(f"Could not find unique name for {result.image_path} after {max_attempts} attempts. Skipping.")
            return
        
        try:
            shutil.copy2(result.image_path, new_image_path)
        except Exception as e:
            logger.error(f"Error copying {result.image_path} to {new_image_path}: {str(e)}")
            return
        
        # Final short name is the new image's base
        final_short_name = os.path.splitext(new_image_name)[0]
        
        # Add to TSV
        self.tsv_handler.upsert_entry(
            result.original_filename,
            final_short_name,
            result.long_desc,
            result.context,
            result.file_hash
        )
    
    def process_all(self) -> None:
        """Process all images in the folder."""
        composite_tasks, skip_filenames = self.collect_composite_tasks()
        tasks = self.collect_image_files(skip_filenames=skip_filenames)
        total_count = len(tasks) + len(composite_tasks)
        composite_count = len(composite_tasks)
        single_image_count = len(tasks)
        composite_image_count = sum(len(files) for _entry, files, _hash in composite_tasks)
        total_images = single_image_count + composite_image_count
        def _label(count: int, singular: str) -> str:
            return singular if count == 1 else f"{singular}s"

        row_label = _label(total_count, "row") if composite_count else _label(total_count, "image")
        
        if total_count == 0:
            self.tsv_handler.write_all()
            logger.info("No new images to process.")
            return

        for filename, _image_path, file_hash, context in tasks:
            self.tsv_handler.upsert_entry(
                orig_filename=filename,
                short_desc="",
                long_desc="",
                context=context,
                file_hash=file_hash
            )
        
        if composite_count:
            logger.info(
                "Processing "
                f"{total_count} {row_label} "
                f"({single_image_count} {_label(single_image_count, 'single image')}, "
                f"{composite_count} {_label(composite_count, 'composite row')}, "
                f"{total_images} {_label(total_images, 'image')} total)..."
            )
        else:
            logger.info(f"Processing {total_count} {row_label}...")

        # We'll use this to track progress
        done_count = 0
        
        try:
            # Parallel describing
            with ThreadPoolExecutor(max_workers=self.max_workers) as executor:
                future_to_kind = {}
                for task in tasks:
                    future_to_kind[executor.submit(self.process_image, task)] = "image"
                for task in composite_tasks:
                    future_to_kind[executor.submit(self.process_composite_task, task)] = "composite"
                
                for future in as_completed(future_to_kind):
                    result = None
                    try:
                        result = future.result()
                    except Exception as e:
                        logger.error(f"Error processing image task: {str(e)}")

                    with self.progress_lock:
                        done_count += 1
                        logger.info(f"{done_count} of {total_count} {row_label}")

                    if result:
                        kind = future_to_kind[future]
                        if kind == "image":
                            self.handle_image_result(result)
                        else:
                            self.handle_composite_result(result)
        finally:
            # Rewrite TSV with updated entries, including successful results before any failure.
            self.tsv_handler.write_all()
        logger.info(f"Processed {done_count} {row_label}.")

    def init_tsv(self, force: bool = False) -> None:
        """Initialize or refresh a TSV with hashes and empty descriptions/context."""
        all_files = sorted(os.listdir(self.folder_path), key=str.lower)
        pending_hashes = set()
        existing_entries = [] if force else list(self.tsv_handler.entries)
        existing_by_hash: Dict[str, TSVEntry] = {}
        existing_singles_by_filename: Dict[str, TSVEntry] = {}
        existing_composites_by_base: Dict[str, TSVEntry] = {}
        matched_existing_ids: Set[int] = set()

        for entry in existing_entries:
            if entry.file_hash and entry.file_hash not in existing_by_hash:
                existing_by_hash[entry.file_hash] = entry
            if entry.composite:
                base_key = self._normalize_base_name(entry.original_filename)
                if base_key and base_key not in existing_composites_by_base:
                    existing_composites_by_base[base_key] = entry
            else:
                filename_key = entry.original_filename.lower()
                if filename_key and filename_key not in existing_singles_by_filename:
                    existing_singles_by_filename[filename_key] = entry

        new_entries: List[TSVEntry] = []
        new_entries_by_hash: Dict[str, TSVEntry] = {}

        def add_entry(entry: TSVEntry) -> None:
            new_entries.append(entry)
            if entry.file_hash and entry.file_hash not in new_entries_by_hash:
                new_entries_by_hash[entry.file_hash] = entry

        def refreshed_entry(
            existing: Optional[TSVEntry],
            orig_filename: str,
            file_hash: str,
            composite: bool,
            preserve_descriptions: bool
        ) -> TSVEntry:
            if existing is not None:
                matched_existing_ids.add(id(existing))
            return TSVEntry(
                original_filename=orig_filename,
                short_desc=existing.short_desc if existing and preserve_descriptions else "",
                long_desc=existing.long_desc if existing and preserve_descriptions else "",
                context=existing.context if existing else "",
                composite=composite,
                file_hash=file_hash
            )

        total_count = 0
        skip_filenames: Set[str] = set()
        if not self.no_composites:
            for base_name, files in self._discover_composite_sets():
                for _index, filename, _path, _file_hash in files:
                    skip_filenames.add(filename.lower())
                composite_hash = self._compute_composite_hash(files)
                preferred_name = self._preferred_composite_name(base_name, files)
                existing = existing_by_hash.get(composite_hash)
                preserve_descriptions = existing is not None
                if existing is None:
                    existing = existing_composites_by_base.get(self._normalize_base_name(base_name))
                    if existing is None:
                        existing = existing_composites_by_base.get(self._normalize_base_name(preferred_name))
                total_count += 1
                add_entry(
                    refreshed_entry(
                        existing=existing,
                        orig_filename=preferred_name,
                        file_hash=composite_hash,
                        composite=True,
                        preserve_descriptions=preserve_descriptions
                    )
                )

        for filename in all_files:
            if not filename.lower().endswith(VALID_EXTENSIONS):
                continue
            if filename.lower() in skip_filenames:
                continue
            image_path = os.path.join(self.folder_path, filename)
            file_hash = FileHelper.hash_file(image_path)
            if file_hash in pending_hashes:
                continue
            pending_hashes.add(file_hash)
            existing = existing_by_hash.get(file_hash)
            preserve_descriptions = existing is not None
            if existing is None:
                existing = existing_singles_by_filename.get(filename.lower())
            total_count += 1
            add_entry(
                refreshed_entry(
                    existing=existing,
                    orig_filename=filename,
                    file_hash=file_hash,
                    composite=False,
                    preserve_descriptions=preserve_descriptions
                )
            )

        if total_count == 0:
            logger.info("No images found to initialize.")
            return

        preserved_count = 0
        if not force:
            for entry in existing_entries:
                if id(entry) in matched_existing_ids:
                    continue
                preserved_count += 1
                add_entry(entry)

        self.tsv_handler.entries = new_entries
        self.tsv_handler.entries_by_hash = new_entries_by_hash
        self.tsv_handler.write_all()
        logger.info(f"Initialized TSV for {total_count} rows.")
        if preserved_count:
            logger.info(f"Preserved {preserved_count} existing unmatched rows.")

    def make_excel(self) -> None:
        """Generate an Excel file from the current TSV entries."""
        if not os.path.exists(self.tsv_path):
            logger.error(f"No TSV found at {self.tsv_path}")
            return
        xlsx_path = str(Path(self.tsv_path).with_suffix(".xlsx"))
        if self.tsv_handler.write_excel(xlsx_path):
            logger.info(f"Wrote Excel file to {xlsx_path}")


class CLI:
    """Command-line interface handler."""
    
    @staticmethod
    def parse_args() -> argparse.Namespace:
        """Parse command-line arguments."""
        parser = argparse.ArgumentParser(
            description="Describe images with OpenAI. For folders, writes 6 columns in TSV:\n"
                        "OriginalFilename, ShortDescription, LongDescription, Context, Composite, SHA1.\n"
                        "For single image files, outputs descriptions directly."
        )
        parser.add_argument("path", nargs="?", help="Path to folder or single image file.")
        parser.add_argument(
            "-t", "--temperature",
            type=float,
            help="Sampling temperature for OpenAI (default=1.0)."
        )
        parser.add_argument(
            "-l", "--length",
            type=int,
            help="Max tokens (default=4000)."
        )
        parser.add_argument(
            "-n", "--no-copy",
            action="store_true",
            help="If provided, do NOT copy files to the output folder (folder mode only)."
        )
        parser.add_argument(
            "-k", "--api-key",
            type=str,
            help="OpenAI API key (overrides config file and environment variable)."
        )
        parser.add_argument(
            "-w", "--workers",
            type=int,
            help="Maximum number of concurrent workers in folder mode (0=auto, default=0)."
        )
        parser.add_argument(
            "-v", "--verbose",
            action="store_true",
            help="Enable verbose output including HTTP requests."
        )
        parser.add_argument(
            "-c", "--config",
            type=str,
            help="Path to the configuration file (default: config.json in the target folder, current directory, or ~/.config/gid/config.json)"
        )
        parser.add_argument(
            "--init-tsv",
            action="store_true",
            help="Generate TSV with hashes and empty descriptions/context (folder mode only; use --composites to include composite rows)."
        )
        parser.add_argument(
            "--force-init-tsv",
            action="store_true",
            help="Reset the TSV when using --init-tsv instead of preserving existing rows/context."
        )
        parser.add_argument(
            "--make-excel",
            action="store_true",
            help="Generate an Excel .xlsx file from the existing TSV (folder mode only)."
        )
        composite_group = parser.add_mutually_exclusive_group()
        composite_group.add_argument(
            "--composites",
            action="store_true",
            help="Enable automatic composite detection (folder mode only)."
        )
        composite_group.add_argument(
            "--no-composites",
            action="store_true",
            help="Disable automatic composite detection (default; useful to override config)."
        )
        parser.add_argument(
            "--show-composites",
            action="store_true",
            help="List discovered composite sets and their matching files (folder mode only; no API calls or output writes)."
        )
        parser.add_argument(
            "--write-sample-config",
            nargs="?",
            const="config.json.sample",
            metavar="PATH",
            help="Write built-in defaults to a sample config file and exit (default: config.json.sample)."
        )
        parser.add_argument(
            "-m", "--model",
            type=str,
            help=f"OpenAI model ID to send to the API (default: {DEFAULT_MODEL})."
        )
        parser.add_argument(
            "-p", "--prompt",
            type=str,
            metavar="NAME",
            help="System prompt file name from a prompts/ directory (for example: web for prompts/web.md)."
        )
        reasoning_group = parser.add_mutually_exclusive_group()
        reasoning_group.add_argument(
            "--reasoning-effort",
            choices=REASONING_EFFORT_VALUES,
            help=(
                f"Reasoning effort for supported models (default: {DEFAULT_REASONING_EFFORT}). "
                f"Choices: {', '.join(REASONING_EFFORT_VALUES)}."
            )
        )
        reasoning_group.add_argument(
            "--no-reasoning",
            action="store_true",
            help="Do not send a reasoning parameter; use this for models that do not support reasoning."
        )
        
        # If user didn't supply any arguments, show help and exit
        if len(sys.argv) == 1:
            parser.print_help(sys.stderr)
            sys.exit(1)
        
        args = parser.parse_args()
        if args.force_init_tsv and not args.init_tsv:
            parser.error("--force-init-tsv requires --init-tsv")
        if args.write_sample_config and (
            args.init_tsv or args.make_excel or args.show_composites
        ):
            parser.error("--write-sample-config cannot be combined with folder no-API actions")
        if args.length is not None and args.length < 1:
            parser.error("--length must be at least 1")
        if args.show_composites and args.no_composites:
            parser.error("--show-composites cannot be combined with --no-composites")
        if args.workers is not None and args.workers < 0:
            parser.error("--workers must be 0 for auto or at least 1")
        if args.temperature is not None and args.temperature < 0:
            parser.error("--temperature must be non-negative")
        return args
    
    @staticmethod
    def get_config(args: argparse.Namespace) -> Dict[str, Any]:
        """Load config from file and override with command line args."""
        # Load config from file
        config_path = args.config
        target_dir = None
        if args.path:
            if os.path.isdir(args.path):
                target_dir = args.path
            elif os.path.isfile(args.path):
                target_dir = os.path.dirname(args.path)
        if not config_path and target_dir:
            candidate = os.path.join(target_dir, "config.json")
            if os.path.exists(candidate):
                config_path = candidate
        if not config_path:
            config_path = Config.find_config_file()
        try:
            config = Config.load_config(config_path, require_exists=bool(args.config))
        except (FileNotFoundError, ValueError) as e:
            print(f"Error: {str(e)}", file=sys.stderr)
            sys.exit(1)

        config.setdefault("api", {})
        config.setdefault("parameters", {})
        config.setdefault("processing", {})
        config.setdefault("output", {})
        config.setdefault("prompt", {})
        config["_gid"] = {
            "prompt_dirs": Config.prompt_dirs(target_dir, config_path),
            "required_prompt_reference_fields": []
        }

        # Override with command line args
        if args.api_key:
            config["api"]["api_key"] = args.api_key
        if args.model:
            config["api"]["model"] = args.model
        if args.prompt:
            config["prompt"]["system_prompt"] = args.prompt
            config["_gid"]["required_prompt_reference_fields"].append("system_prompt")
        if args.reasoning_effort:
            config["parameters"]["reasoning_effort"] = args.reasoning_effort
        if args.no_reasoning:
            config["parameters"]["reasoning_effort"] = None
        if args.temperature is not None:
            config["parameters"]["temperature"] = args.temperature
        if args.length is not None:
            config["parameters"]["max_tokens"] = args.length
        if args.no_copy:
            config["processing"]["no_copy"] = True
        if args.composites or args.show_composites:
            config["processing"]["no_composites"] = False
        if args.no_composites:
            config["processing"]["no_composites"] = True
        if args.workers is not None:
            config["processing"]["max_workers"] = args.workers
        if args.verbose:
            config["processing"]["verbose"] = True

        # Check for API key in environment if not in config or args
        if config["api"].get("api_key") == "...":
            config["api"]["api_key"] = ""
        if not config["api"].get("api_key"):
            config["api"]["api_key"] = os.environ.get("OPENAI_API_KEY", "")

        try:
            config["api"]["model"] = Config.normalize_model_name(config["api"].get("model"))
            config["parameters"]["reasoning_effort"] = Config.normalize_reasoning_effort(
                config["parameters"].get("reasoning_effort")
            )
            Config.normalize_max_workers(config["processing"].get("max_workers"))
        except (FileNotFoundError, ValueError) as e:
            print(f"Error: {str(e)}", file=sys.stderr)
            sys.exit(1)

        return config
    
    @staticmethod
    def process_single_image(image_path: str, config: Dict[str, Any]) -> None:
        """Process a single image file and output descriptions directly."""
        Config.validate_prompt_config(config)
        if config["processing"]["verbose"]:
            # Enable HTTP request logging in verbose mode
            logging.getLogger("openai").setLevel(logging.INFO)
            logging.getLogger("httpx").setLevel(logging.INFO)
            
        describer = ImageDescriber(
            api_key=config["api"]["api_key"],
            model=config["api"]["model"],
            temperature=config["parameters"]["temperature"],
            max_tokens=config["parameters"]["max_tokens"],
            reasoning_effort=config["parameters"].get("reasoning_effort"),
            system_prompt=config["prompt"]["system_prompt"],
            instructions_prompt=config["prompt"]["instructions_prompt"],
            single_image_prompt=config["prompt"]["single_image_prompt"],
            composite_image_prompt=config["prompt"]["composite_image_prompt"],
            context_template=config["prompt"]["context_template"],
            short_description_max_words=config["prompt"]["short_description_max_words"]
        )
        
        short_desc, long_desc = describer.describe_image(image_path)
        
        if short_desc == "Error":
            print(f"Error: {long_desc}")
            sys.exit(1)
            
        print(short_desc)
        print(long_desc)
    
    @staticmethod
    def run() -> None:
        """Run the command-line interface."""
        args = CLI.parse_args()

        if args.write_sample_config:
            Config.write_sample_config(args.write_sample_config)
            logger.info(f"Wrote sample config to {args.write_sample_config}")
            return

        # If user didn't supply the path argument, show usage and exit
        if not args.path:
            print("Error: Path to folder or image file is required.", file=sys.stderr)
            sys.exit(1)
        
        # Get configuration
        config = CLI.get_config(args)
        
        # Check for API key unless we're in a no-API mode
        if not args.init_tsv and not args.show_composites and not args.make_excel and not config["api"].get("api_key"):
            print("Error: OpenAI API key not provided. Use -k/--api-key, set OPENAI_API_KEY environment variable, or add it to config.json.", file=sys.stderr)
            sys.exit(1)
        
        try:
            # Check if path is a directory or a file
            if os.path.isdir(args.path):
                if args.show_composites or args.init_tsv or args.make_excel:
                    processor = ImageProcessor(
                        folder_path=args.path,
                        config=config,
                        init_only=True
                    )
                    if args.show_composites:
                        processor.show_composites()
                    if args.init_tsv:
                        processor.init_tsv(force=args.force_init_tsv)
                    if args.make_excel:
                        processor.make_excel()
                else:
                    # Process folder
                    processor = ImageProcessor(
                        folder_path=args.path,
                        config=config
                    )
                    processor.process_all()
            elif os.path.isfile(args.path) and args.path.lower().endswith(VALID_EXTENSIONS):
                if args.init_tsv:
                    print("Error: --init-tsv is only supported for folder mode.", file=sys.stderr)
                    sys.exit(1)
                if args.show_composites:
                    print("Error: --show-composites is only supported for folder mode.", file=sys.stderr)
                    sys.exit(1)
                if args.make_excel:
                    print("Error: --make-excel is only supported for folder mode.", file=sys.stderr)
                    sys.exit(1)
                # Process single image
                CLI.process_single_image(
                    image_path=args.path,
                    config=config
                )
            else:
                print(f"Error: '{args.path}' is not a valid directory or image file.", file=sys.stderr)
                sys.exit(1)
        except (FileNotFoundError, ValueError) as e:
            print(f"Error: {str(e)}", file=sys.stderr)
            sys.exit(1)


def main():
    """Main entry point."""
    CLI.run()


if __name__ == "__main__":
    main()
